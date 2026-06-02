import io
from pptx import Presentation as PptxPresentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

MARGIN = Inches(0.8)
TITLE_H = Inches(1.0)
CONTENT_TOP = Inches(1.7)
CONTENT_H = SLIDE_H - CONTENT_TOP - MARGIN

COLOR_HEADING = RGBColor(0x11, 0x18, 0x27)
COLOR_BODY = RGBColor(0x37, 0x41, 0x51)
COLOR_MUTED = RGBColor(0x6B, 0x72, 0x80)


def _add_text(slide, left, top, width, height, text: str, size_pt: float,
              bold=False, italic=False, color=COLOR_BODY, wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def _add_bullets(slide, left, top, width, height, items: list[str], size_pt: float = 18):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = "•  " + item
        run.font.size = Pt(size_pt)
        run.font.color.rgb = COLOR_BODY
        if i > 0:
            p.space_before = Pt(6)
    return box


def _set_notes(slide, text: str):
    if not text:
        return
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = text


def _blank_slide(prs: PptxPresentation):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank


def build_pptx(presentation_title: str, slides: list[dict]) -> bytes:
    prs = PptxPresentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    content_w = SLIDE_W - 2 * MARGIN

    for slide_data in slides:
        layout = slide_data.get("layout", "bullets")
        title = slide_data.get("title", "")
        content = slide_data.get("content") or {}
        notes = slide_data.get("speaker_notes", "")

        slide = _blank_slide(prs)

        if layout == "title":
            heading = content.get("heading") or title
            subheading = content.get("subheading", "")
            _add_text(slide, MARGIN, Inches(2.2), content_w, Inches(1.8),
                      heading, 40, bold=True, color=COLOR_HEADING)
            if subheading:
                _add_text(slide, MARGIN, Inches(4.2), content_w, Inches(0.9),
                          subheading, 22, color=COLOR_MUTED)

        elif layout == "bullets":
            bullets = content.get("bullets") or []
            _add_text(slide, MARGIN, MARGIN, content_w, TITLE_H,
                      title, 28, bold=True, color=COLOR_HEADING)
            if bullets:
                _add_bullets(slide, MARGIN, CONTENT_TOP, content_w, CONTENT_H, bullets)

        elif layout == "quote":
            quote = content.get("quote", "")
            attribution = content.get("attribution", "")
            _add_text(slide, MARGIN, MARGIN, content_w, TITLE_H,
                      title, 28, bold=True, color=COLOR_HEADING)
            if quote:
                box = slide.shapes.add_textbox(MARGIN, Inches(2.0), content_w, Inches(3.5))
                tf = box.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                run = p.add_run()
                run.text = f"“{quote}”"
                run.font.size = Pt(24)
                run.font.italic = True
                run.font.color.rgb = COLOR_HEADING
                if attribution:
                    p2 = tf.add_paragraph()
                    p2.space_before = Pt(14)
                    run2 = p2.add_run()
                    run2.text = f"— {attribution}"
                    run2.font.size = Pt(16)
                    run2.font.color.rgb = COLOR_MUTED

        elif layout == "two-col":
            left_items = content.get("left") or []
            right_items = content.get("right") or []
            col_w = (content_w - Inches(0.5)) / 2
            _add_text(slide, MARGIN, MARGIN, content_w, TITLE_H,
                      title, 28, bold=True, color=COLOR_HEADING)
            if left_items:
                _add_bullets(slide, MARGIN, CONTENT_TOP, col_w, CONTENT_H, left_items, 16)
            if right_items:
                _add_bullets(slide, MARGIN + col_w + Inches(0.5), CONTENT_TOP,
                             col_w, CONTENT_H, right_items, 16)

        elif layout == "embed":
            url = content.get("url", "")
            caption = content.get("caption", "")
            _add_text(slide, MARGIN, MARGIN, content_w, TITLE_H,
                      title, 28, bold=True, color=COLOR_HEADING)
            body = f"[{url}]" if url else ""
            if caption:
                body = (body + "\n\n" + caption).strip()
            if body:
                _add_text(slide, MARGIN, CONTENT_TOP, content_w, CONTENT_H,
                          body, 16, color=COLOR_MUTED)

        _set_notes(slide, notes)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
